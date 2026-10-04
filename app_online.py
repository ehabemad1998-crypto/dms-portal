import streamlit as st
import requests
import os
import re
import io
import uuid
import time
import html
import base64
import zipfile
from datetime import datetime
from urllib.parse import quote
from zoneinfo import ZoneInfo

# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="Document Control Portal",
    page_icon="📁",
    layout="wide"
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

BUCKET = "project-files"


# =========================================================
# SECRETS (set in Streamlit Cloud -> App settings -> Secrets)
# =========================================================

def secret(name, default=""):

    try:
        value = st.secrets[name]
    except Exception:
        value = os.environ.get(name, default)

    return str(value).strip()


SUPABASE_URL = secret("SUPABASE_URL").rstrip("/")

SUPABASE_KEY = secret("SUPABASE_ANON_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:

    st.error(
        "Missing SUPABASE_URL or SUPABASE_ANON_KEY. "
        "Add them in the app Secrets."
    )

    st.stop()


# =========================================================
# SUPABASE LOW-LEVEL (plain HTTPS, no extra packages)
# =========================================================

def get_token():
    """Return a valid login token, refreshing it when it is about to expire."""

    auth = st.session_state.get("auth")

    if not auth:
        return None

    if time.time() > auth["expires_at"] - 60:

        try:

            r = requests.post(
                f"{SUPABASE_URL}/auth/v1/token",
                params={"grant_type": "refresh_token"},
                json={"refresh_token": auth["refresh"]},
                headers={"apikey": SUPABASE_KEY},
                timeout=30
            )

            if r.status_code == 200:

                d = r.json()

                auth["token"] = d["access_token"]
                auth["refresh"] = d["refresh_token"]
                auth["expires_at"] = time.time() + d.get("expires_in", 3600)

            else:

                st.session_state.pop("auth", None)

                return None

        except Exception:

            return auth["token"]

    return auth["token"]


def api(method, path, params=None, json_body=None, data=None,
        headers=None, timeout=60):

    h = {"apikey": SUPABASE_KEY}

    token = get_token()

    if token:
        h["Authorization"] = f"Bearer {token}"

    if headers:
        h.update(headers)

    return requests.request(
        method,
        f"{SUPABASE_URL}{path}",
        params=params,
        json=json_body,
        data=data,
        headers=h,
        timeout=timeout
    )


def err(r):
    return f"{r.status_code}: {r.text[:200]}"


def login(email, password):

    try:

        r = requests.post(
            f"{SUPABASE_URL}/auth/v1/token",
            params={"grant_type": "password"},
            json={"email": email, "password": password},
            headers={"apikey": SUPABASE_KEY},
            timeout=30
        )

    except Exception as e:

        return False, f"Connection error: {e}"

    if r.status_code in (400, 401):
        return False, "Wrong email or password."

    if r.status_code != 200:

        return False, (
            f"Service error ({r.status_code}). "
            "The database may be paused or the keys are wrong."
        )

    d = r.json()

    auth = {
        "token": d["access_token"],
        "refresh": d["refresh_token"],
        "expires_at": time.time() + d.get("expires_in", 3600),
        "email": email,
        "user_id": d["user"]["id"],
        "role": "viewer"
    }

    st.session_state.auth = auth

    rr = api(
        "GET",
        "/rest/v1/profiles",
        params={"select": "role", "id": f"eq.{auth['user_id']}"}
    )

    if rr.status_code == 200 and rr.json():
        auth["role"] = rr.json()[0]["role"]

    return True, ""


# =========================================================
# HELPERS
# =========================================================

def fmt_time(value):

    try:

        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))

        return dt.astimezone(
            ZoneInfo("Africa/Cairo")
        ).strftime("%Y-%m-%d %H:%M")

    except Exception:

        return str(value)[:16].replace("T", " ")


def format_size(size):

    size = float(size or 0)

    for unit in ["B", "KB", "MB", "GB"]:

        if size < 1024:
            return f"{size:.1f} {unit}"

        size /= 1024

    return f"{size:.1f} TB"


def number_key(number):
    """Sort 1, 2, 10 as numbers, then any non-numeric ones as text."""

    if str(number).isdigit():
        return (0, int(number), str(number))

    return (1, 0, str(number))


def show_table(rows):
    """Simple HTML table (no pyarrow needed)."""

    if not rows:
        return

    headers = list(rows[0].keys())

    head = "".join(
        f"<th style='text-align:left;padding:8px;"
        f"border-bottom:2px solid rgba(128,128,128,0.4)'>"
        f"{html.escape(h)}</th>"
        for h in headers
    )

    body = ""

    for r in rows:

        cells = "".join(
            f"<td style='padding:8px;"
            f"border-bottom:1px solid rgba(128,128,128,0.25)'>"
            f"{html.escape(str(r[h] if r[h] is not None else ''))}</td>"
            for h in headers
        )

        body += f"<tr>{cells}</tr>"

    st.markdown(
        "<div style='overflow-x:auto'>"
        "<table style='width:100%;border-collapse:collapse;font-size:14px'>"
        f"<thead><tr>{head}</tr></thead><tbody>{body}</tbody>"
        "</table></div>",
        unsafe_allow_html=True
    )


# =========================================================
# PROJECTS (database)
# =========================================================

def get_projects(recent=False):
    """List of (number, name, client, status). recent=True: newest first."""

    params = {
        "select": "id,project_number,project_name,client,status",
        "order": "id.desc",
        "limit": 5000
    }

    r = api("GET", "/rest/v1/projects", params=params)

    if r.status_code != 200:

        st.error(f"Could not load projects ({err(r)})")

        return []

    rows = r.json()

    if not recent:
        rows.sort(key=lambda x: number_key(x["project_number"]))

    return [
        (
            x["project_number"],
            x["project_name"],
            x["client"] or "",
            x["status"]
        )
        for x in rows
    ]


def get_project_id(number):

    r = api(
        "GET",
        "/rest/v1/projects",
        params={"select": "id", "project_number": f"eq.{number}"}
    )

    if r.status_code == 200 and r.json():
        return r.json()[0]["id"]

    return None


def add_project(number, name, client, status):

    r = api(
        "POST",
        "/rest/v1/projects",
        json_body={
            "project_number": number,
            "project_name": name,
            "client": client,
            "status": status
        },
        headers={"Prefer": "return=minimal"}
    )

    if r.status_code in (200, 201, 204):
        return True, "Project added successfully."

    if r.status_code == 409:
        return False, "Project Number already exists."

    if r.status_code in (401, 403):
        return False, "You are not allowed to add projects."

    return False, f"Could not add the project ({err(r)})"


def bulk_add_projects(rows, status):
    """Insert many projects. Existing numbers are skipped."""

    added = 0

    for i in range(0, len(rows), 200):

        chunk = [
            {
                "project_number": n,
                "project_name": name,
                "client": client,
                "status": status
            }
            for n, name, client in rows[i:i + 200]
        ]

        r = api(
            "POST",
            "/rest/v1/projects",
            params={"on_conflict": "project_number"},
            json_body=chunk,
            headers={
                "Prefer": "resolution=ignore-duplicates,"
                          "return=representation"
            }
        )

        if r.status_code not in (200, 201):
            return None, None, f"Import failed ({err(r)})"

        added += len(r.json())

    return added, len(rows) - added, None


def delete_storage(paths):

    if paths:

        api(
            "DELETE",
            f"/storage/v1/object/{BUCKET}",
            json_body={"prefixes": paths}
        )


def delete_project(number):
    """Delete the project, its file records, AI results and stored files."""

    paths = [d[7] for d in get_documents(number)]

    r = api(
        "DELETE",
        "/rest/v1/projects",
        params={"project_number": f"eq.{number}"},
        headers={"Prefer": "return=representation"}
    )

    if r.status_code not in (200, 204):
        return False, f"Could not delete ({err(r)})"

    if not r.json():
        return False, "Nothing was deleted (not allowed, or project not found)."

    delete_storage(paths)

    return True, f"Project {number} was deleted with its files."


def read_projects_excel(file):
    """Read projects from an Excel file (first sheet, header in row 1).
    Returns (rows, error). Each row is (number, name, client)."""

    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "The package openpyxl is missing."

    try:
        wb = load_workbook(file, data_only=True, read_only=True)
        data = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return None, f"Could not read the file: {e}"

    if not data:
        return None, "The file is empty."

    header = [str(h or "").strip().lower() for h in data[0]]

    def find(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_client = find("client name", "client")
    i_name = find("project name", "name")
    i_number = find("project number", "project no", "project no.", "number")

    if i_name is None or i_number is None:
        return None, (
            "Could not find the columns 'Project Name' and "
            "'Project Number' in the first row."
        )

    def clean(v):

        if v is None:
            return ""

        if isinstance(v, float) and v.is_integer():
            v = int(v)

        return " ".join(str(v).split())

    def cell(row, i):
        return clean(row[i]) if i is not None and i < len(row) else ""

    rows = []

    for r in data[1:]:

        number = cell(r, i_number)
        name = cell(r, i_name)
        client = cell(r, i_client)

        if number and name:
            rows.append((number, name, client))

    if not rows:
        return None, "No projects found in the file."

    return rows, None


# =========================================================
# DOCUMENTS (database + file storage)
# =========================================================

def add_document(project_number, file_name, storage_path,
                 file_type, file_size, notes):

    r = api(
        "POST",
        "/rest/v1/documents",
        json_body={
            "project_number": project_number,
            "file_name": file_name,
            "storage_path": storage_path,
            "file_type": file_type,
            "file_size": file_size,
            "notes": notes
        },
        headers={"Prefer": "return=minimal"}
    )

    if r.status_code in (200, 201, 204):
        return True, ""

    return False, f"Could not save the file record ({err(r)})"


def get_documents(project_number=None):
    """List of (id, project, name, type, size, notes, uploaded, path)."""

    params = {
        "select": "id,project_number,file_name,file_type,file_size,"
                  "notes,uploaded_at,storage_path",
        "order": "id.desc",
        "limit": 5000
    }

    if project_number:
        params["project_number"] = f"eq.{project_number}"

    r = api("GET", "/rest/v1/documents", params=params)

    if r.status_code != 200:

        st.error(f"Could not load documents ({err(r)})")

        return []

    return [
        (
            d["id"],
            d["project_number"],
            d["file_name"],
            d["file_type"],
            d["file_size"],
            d["notes"] or "",
            fmt_time(d["uploaded_at"]),
            d["storage_path"]
        )
        for d in r.json()
    ]


def count_documents():

    r = api(
        "GET",
        "/rest/v1/documents",
        params={"select": "id", "limit": 1},
        headers={"Prefer": "count=exact"}
    )

    try:
        return int(r.headers.get("Content-Range", "*/0").split("/")[-1])
    except Exception:
        return 0


def upload_file(project_id, file_name, data, content_type):
    """Store a file in the private bucket. Returns (path, error)."""

    ext = re.sub(
        r"[^a-z0-9]", "",
        os.path.splitext(file_name)[1].lower()
    )

    path = f"p{project_id}/{uuid.uuid4().hex}"

    if ext:
        path += f".{ext}"

    r = api(
        "POST",
        f"/storage/v1/object/{BUCKET}/{path}",
        data=data,
        headers={"Content-Type": content_type or "application/octet-stream"},
        timeout=300
    )

    if r.status_code in (200, 201):
        return path, None

    return None, f"Upload failed ({err(r)})"


def signed_url(path, file_name):
    """Temporary (1 hour) download link for a stored file."""

    r = api(
        "POST",
        f"/storage/v1/object/sign/{BUCKET}/{path}",
        json_body={"expiresIn": 3600}
    )

    if r.status_code != 200:
        return None

    body = r.json()

    signed = body.get("signedURL") or body.get("signedUrl")

    if not signed:
        return None

    if signed.startswith("http"):
        url = signed

    elif signed.startswith("/storage/v1"):
        url = f"{SUPABASE_URL}{signed}"

    else:

        if not signed.startswith("/"):
            signed = "/" + signed

        url = f"{SUPABASE_URL}/storage/v1{signed}"

    joiner = "&" if "?" in url else "?"

    return f"{url}{joiner}download={quote(file_name)}"


def download_bytes(path):

    r = api(
        "GET",
        f"/storage/v1/object/authenticated/{BUCKET}/{path}",
        timeout=300
    )

    return r.content if r.status_code == 200 else None


# =========================================================
# AI
# =========================================================

AI_MODEL = "claude-sonnet-5-5"

AI_URL = "https://api.anthropic.com/v1/messages"

AI_MAX_CHARS = 60000

AI_DEFAULT_INSTRUCTION = (
    "Summarize this file for a document controller: what it is, "
    "the key data (numbers, dates, revisions, quantities), "
    "and anything missing or unclear."
)

AI_SYSTEM_PROMPT = (
    "You are an assistant for a document controller at an engineering "
    "company. Be accurate and concise. Never invent numbers or data that "
    "are not in the file; say clearly when something is missing or unclear. "
    "Reply in the same language as the user's instruction."
)

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp"
}


def save_ai_result(project_number, file_name, request_text, result):

    api(
        "POST",
        "/rest/v1/ai_results",
        json_body={
            "project_number": project_number,
            "file_name": file_name,
            "request": request_text,
            "result": result
        },
        headers={"Prefer": "return=minimal"}
    )


def get_ai_results(project_number):
    """List of (file_name, request, result, created_at)."""

    r = api(
        "GET",
        "/rest/v1/ai_results",
        params={
            "select": "file_name,request,result,created_at",
            "project_number": f"eq.{project_number}",
            "order": "id.desc"
        }
    )

    if r.status_code != 200:
        return []

    return [
        (
            x["file_name"],
            x["request"] or "",
            x["result"] or "",
            fmt_time(x["created_at"])
        )
        for x in r.json()
    ]


def excel_to_text(data, max_rows=300, max_cols=40):

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)

    parts = []

    for ws in wb.worksheets:

        parts.append(f"=== Sheet: {ws.title} ===")

        for i, row in enumerate(ws.iter_rows(values_only=True)):

            if i >= max_rows:
                parts.append(f"... (cut after {max_rows} rows)")
                break

            cells = [
                "" if c is None else str(c)
                for c in row[:max_cols]
            ]

            if any(cells):
                parts.append(" | ".join(cells))

    wb.close()

    return "\n".join(parts)[:AI_MAX_CHARS]


def docx_to_text(data):

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")

    xml = xml.replace("</w:p>", "\n")

    text = re.sub(r"<[^>]+>", "", xml)

    return html.unescape(text)[:AI_MAX_CHARS]


def build_file_block(name, data):
    """Turn a file into a message block for the AI. Returns (block, error)."""

    ext = os.path.splitext(name)[1].lower()

    try:

        if ext == ".pdf" or ext in IMAGE_TYPES:

            if len(data) > 20 * 1024 * 1024:
                return None, f"{name} is larger than 20 MB, too big for AI."

            encoded = base64.b64encode(data).decode()

            if ext == ".pdf":

                return {
                    "type": "document",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": encoded
                    }
                }, None

            return {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": IMAGE_TYPES[ext],
                    "data": encoded
                }
            }, None

        if ext in (".xlsx", ".xlsm"):
            text = excel_to_text(data)

        elif ext == ".docx":
            text = docx_to_text(data)

        elif ext in (".csv", ".txt", ".md", ".json"):
            text = data.decode("utf-8", errors="replace")[:AI_MAX_CHARS]

        else:

            return None, (
                f"File type {ext or '(none)'} is not supported by AI yet. "
                "Supported: Excel (.xlsx), PDF, Word (.docx), "
                "images, CSV/TXT."
            )

    except Exception as e:
        return None, f"Could not read {name}: {e}"

    return {
        "type": "text",
        "text": f"File name: {name}\n\n{text}"
    }, None


def analyze_file(name, data, instruction):
    """Send one file to the AI. Returns (ok, text_or_error)."""

    key = secret("ANTHROPIC_API_KEY")

    if not key:
        return False, "No ANTHROPIC_API_KEY found in the app Secrets."

    block, error = build_file_block(name, data)

    if error:
        return False, error

    payload = {
        "model": AI_MODEL,
        "max_tokens": 2000,
        "system": AI_SYSTEM_PROMPT,
        "messages": [
            {
                "role": "user",
                "content": [block, {"type": "text", "text": instruction}]
            }
        ]
    }

    try:

        response = requests.post(
            AI_URL,
            headers={
                "x-api-key": key,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json"
            },
            json=payload,
            timeout=180
        )

    except Exception as e:

        return False, f"Could not connect to the AI service: {e}"

    if response.status_code != 200:

        return False, (
            f"AI service error {response.status_code}: "
            f"{response.text[:300]}"
        )

    text = "\n".join(
        b.get("text", "")
        for b in response.json().get("content", [])
        if b.get("type") == "text"
    ).strip()

    return True, text or "(AI returned an empty answer.)"


# ===== UI =====

# =========================================================
# LOGO
# =========================================================

logo = None

logo_path = os.path.join(BASE_DIR, "logo.png")

if os.path.exists(logo_path):

    with open(logo_path, "rb") as image_file:
        logo = base64.b64encode(image_file.read()).decode()

if logo:

    st.markdown(
        f"""
        <div style="width:100%;text-align:center;
                    margin-top:-35px;margin-bottom:15px;">
            <img src="data:image/png;base64,{logo}"
                 style="width:200px;height:auto;">
        </div>
        """,
        unsafe_allow_html=True
    )


# =========================================================
# LOGIN
# =========================================================

if not get_token():

    st.title("📁 Document Control Portal")

    st.write("Please sign in to continue.")

    with st.form("login_form"):

        email = st.text_input("Email")

        password = st.text_input("Password", type="password")

        submitted = st.form_submit_button(
            "Sign in",
            use_container_width=True
        )

    if submitted:

        if not email or not password:

            st.error("Please enter your email and password.")

        else:

            ok, message = login(email.strip(), password)

            if ok:
                st.rerun()
            else:
                st.error(message)

    st.stop()


auth = st.session_state["auth"]

is_admin = auth["role"] == "admin"


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title("📁 ECE")

st.sidebar.caption("Document Control Portal")

st.sidebar.divider()

menu = [
    "🏠 Dashboard",
    "📁 Projects",
    "📄 Documents"
]

if is_admin:
    menu.append("📤 Upload")

menu += ["📊 Reports", "🪄 AI"]

page = st.sidebar.radio("Menu", menu)

st.sidebar.divider()

st.sidebar.caption(
    f"👤 {auth['email']}  \n"
    f"Role: **{'Admin' if is_admin else 'Viewer (read only)'}**"
)

if st.sidebar.button("Sign out", use_container_width=True):

    st.session_state.pop("auth", None)

    st.rerun()


# =========================================================
# DASHBOARD
# =========================================================

if page == "🏠 Dashboard":

    st.title("📊 Document Control Portal")

    st.write("Welcome to your Document Management System")

    projects = get_projects(recent=True)

    col1, col2, col3, col4 = st.columns(4)

    col1.metric("Projects", len(projects))

    col2.metric("Documents", count_documents())

    col3.metric("Pending", "0")

    col4.metric("Approved", "0")

    st.divider()

    st.subheader("Recent Projects")

    if projects:

        for project in projects[:5]:

            st.write(
                f"📁 **{project[0]}** — "
                f"{project[1]} — "
                f"{project[3]}"
            )

    else:

        st.info("No projects have been added yet.")


# =========================================================
# PROJECTS
# =========================================================

elif page == "📁 Projects":

    st.title("📁 Projects")

    st.write("Manage and track all project information.")

    for key_name in ("import_msg", "delete_msg"):

        if key_name in st.session_state:
            st.success(st.session_state.pop(key_name))

    st.divider()

    search = st.text_input(
        "🔍 Search Project",
        placeholder="Enter project number or project name..."
    )

    projects = get_projects()

    if search:

        projects = [
            p for p in projects
            if search.lower() in p[0].lower()
            or search.lower() in p[1].lower()
        ]

    if projects:

        for project in projects:

            with st.container(border=True):

                col1, col2, col3, col4 = st.columns([1.5, 3, 2, 1.5])

                col1.write(f"**{project[0]}**")

                col2.write(f"**{project[1]}**")

                col3.write(project[2])

                col4.write(
                    "🟢 Active" if project[3] == "Active" else "🔴 Closed"
                )

    else:

        st.info("No projects found.")

    # -----------------------------------------------------
    # ADMIN ONLY: add / import / delete
    # -----------------------------------------------------

    if is_admin:

        st.divider()

        st.subheader("➕ Add New Project")

        with st.form("add_project_form"):

            project_number = st.text_input(
                "Project Number",
                placeholder="Example: 0000"
            )

            project_name = st.text_input(
                "Project Name",
                placeholder="Example Project"
            )

            client = st.text_input("Client")

            status = st.selectbox("Status", ["Active", "Closed"])

            submitted = st.form_submit_button(
                "💾 Save Project",
                use_container_width=True
            )

            if submitted:

                if not project_number:

                    st.error("Please enter Project Number.")

                elif not project_name:

                    st.error("Please enter Project Name.")

                else:

                    success, message = add_project(
                        project_number.strip(),
                        project_name.strip(),
                        client.strip(),
                        status
                    )

                    if success:

                        st.session_state.import_msg = message

                        st.rerun()

                    else:

                        st.error(message)

        st.divider()

        with st.expander("📥 Import Projects from Excel"):

            st.write(
                "The first row must have these columns: "
                "**Client Name**, **Project Name**, **Project Number**. "
                "Project numbers that already exist are skipped."
            )

            import_file = st.file_uploader(
                "Choose Excel file (.xlsx)",
                type=["xlsx"],
                key="import_projects_file"
            )

            import_status = st.selectbox(
                "Status for imported projects",
                ["Active", "Closed"],
                key="import_projects_status"
            )

            if import_file is not None:

                import_rows, import_error = read_projects_excel(import_file)

                if import_error:

                    st.error(import_error)

                else:

                    st.info(
                        f"{len(import_rows)} projects found in the file."
                    )

                    if st.button(
                        "📥 Import",
                        key="import_projects_button",
                        use_container_width=True
                    ):

                        added, skipped, problem = bulk_add_projects(
                            import_rows,
                            import_status
                        )

                        if problem:

                            st.error(problem)

                        else:

                            st.session_state.import_msg = (
                                f"Added {added} project(s). "
                                f"Skipped {skipped} (number already exists)."
                            )

                            st.rerun()

        with st.expander("🗑️ Delete Project"):

            all_projects = get_projects()

            if not all_projects:

                st.info("No projects to delete.")

            else:

                if "delete_counter" not in st.session_state:
                    st.session_state.delete_counter = 0

                counter = st.session_state.delete_counter

                del_options = {
                    f"{p[0]} — {p[1]}": p for p in all_projects
                }

                del_choice = st.selectbox(
                    "Select the project to delete",
                    list(del_options.keys()),
                    key=f"delete_select_{counter}"
                )

                del_project = del_options[del_choice]

                del_number = del_project[0]

                file_count = len(get_documents(del_number))

                warning = (
                    f"You are about to delete project {del_number} — "
                    f"{del_project[1]}. This cannot be undone."
                )

                if file_count:
                    warning += (
                        f" Its {file_count} uploaded file(s) and "
                        "AI results will be deleted too."
                    )

                st.warning(warning)

                confirm_delete = st.checkbox(
                    "Yes, I want to delete this project",
                    key=f"delete_confirm_{counter}_{del_number}"
                )

                if st.button(
                    "🗑️ Delete Project",
                    key=f"delete_button_{counter}",
                    disabled=not confirm_delete,
                    use_container_width=True
                ):

                    ok, message = delete_project(del_number)

                    st.session_state.delete_counter += 1

                    st.session_state.delete_msg = message

                    st.rerun()


# =========================================================
# DOCUMENTS
# =========================================================

elif page == "📄 Documents":

    st.title("📄 Documents")

    projects = get_projects()

    if not projects:

        st.info("No projects yet.")

    else:

        options = {"All Projects": None}

        for p in projects:
            options[f"{p[0]} — {p[1]}"] = p[0]

        col1, col2 = st.columns([2, 2])

        choice = col1.selectbox("Project", list(options.keys()))

        search_file = col2.text_input(
            "🔍 Search File Name",
            placeholder="Type part of the file name..."
        )

        documents = get_documents(options[choice])

        if search_file:

            documents = [
                d for d in documents
                if search_file.lower() in d[2].lower()
            ]

        if documents:

            show_table(
                [
                    {
                        "Project": d[1],
                        "File Name": d[2],
                        "Type": d[3],
                        "Size": format_size(d[4]),
                        "Notes": d[5],
                        "Uploaded At": d[6]
                    }
                    for d in documents
                ]
            )

            st.divider()

            st.subheader("⬇️ Download a File")

            file_options = {
                f"{d[1]} | {d[2]} | {d[6]}": d for d in documents
            }

            selected = st.selectbox(
                "Select file",
                list(file_options.keys())
            )

            doc = file_options[selected]

            url = signed_url(doc[7], doc[2])

            if url:

                st.link_button("⬇️ Download", url)

                st.caption("The download link works for one hour.")

            else:

                st.error("Could not create a download link for this file.")

        else:

            st.info("No documents found.")


# =========================================================
# UPLOAD (admin only)
# =========================================================

elif page == "📤 Upload" and is_admin:

    st.title("📤 Upload Documents")

    st.write(
        "Choose the project, then pick any files "
        "(Excel, CAD, PDF, ...). They are stored online under the project."
    )

    projects = get_projects()

    if not projects:

        st.warning("No projects yet. Add a project first.")

    else:

        if "uploader_key" not in st.session_state:
            st.session_state.uploader_key = 0

        key = st.session_state.uploader_key

        options = {f"{p[0]} — {p[1]}": p for p in projects}

        choice = st.selectbox("📁 Select Project", list(options.keys()))

        project = options[choice]

        project_number = project[0]

        files = st.file_uploader(
            "Choose files",
            accept_multiple_files=True,
            key=f"uploader_{key}"
        )

        notes = st.text_input("Notes (optional)", key=f"notes_{key}")

        send_ai = st.checkbox(
            "🪄 Send to AI after saving",
            value=False,
            key=f"send_ai_{key}"
        )

        ai_instruction = AI_DEFAULT_INSTRUCTION

        if send_ai:

            ai_instruction = st.text_area(
                "What should AI do with the file?",
                value=AI_DEFAULT_INSTRUCTION,
                height=100,
                key=f"ai_instr_{key}"
            )

        if st.button("💾 Save to Project", use_container_width=True):

            if not files:

                st.error("Please choose at least one file.")

            else:

                project_id = get_project_id(project_number)

                if project_id is None:

                    st.error("Project not found.")

                else:

                    ai_outputs = []

                    problems = []

                    saved = 0

                    for f in files:

                        data = f.getvalue()

                        with st.spinner(f"Uploading {f.name} ..."):

                            path, error = upload_file(
                                project_id,
                                f.name,
                                data,
                                f.type
                            )

                        if error:

                            problems.append(f"{f.name}: {error}")

                            continue

                        ext = os.path.splitext(f.name)[1] \
                            .replace(".", "").upper()

                        ok, message = add_document(
                            project_number,
                            f.name,
                            path,
                            ext or "FILE",
                            len(data),
                            notes.strip()
                        )

                        if not ok:

                            delete_storage([path])

                            problems.append(f"{f.name}: {message}")

                            continue

                        saved += 1

                        if send_ai:

                            with st.spinner(
                                f"🪄 AI is analyzing {f.name} ..."
                            ):

                                ai_ok, result = analyze_file(
                                    f.name,
                                    data,
                                    ai_instruction
                                )

                            if ai_ok:

                                save_ai_result(
                                    project_number,
                                    f.name,
                                    ai_instruction,
                                    result
                                )

                            ai_outputs.append((f.name, ai_ok, result))

                    st.session_state.ai_outputs = ai_outputs

                    st.session_state.upload_problems = problems

                    st.session_state.upload_msg = (
                        f"{saved} file(s) saved to project "
                        f"{project_number}."
                    )

                    st.session_state.uploader_key += 1

                    st.rerun()

        if "upload_msg" in st.session_state:

            st.success(st.session_state.pop("upload_msg"))

        for problem in st.session_state.pop("upload_problems", []):
            st.error(problem)

        for name, ok, text in st.session_state.pop("ai_outputs", []):

            if ok:

                with st.expander(f"🪄 AI result — {name}", expanded=True):
                    st.markdown(text)

            else:

                st.error(f"🪄 AI could not analyze {name}: {text}")

        st.divider()

        st.subheader("Files in this project")

        documents = get_documents(project_number)

        if documents:

            show_table(
                [
                    {
                        "File Name": d[2],
                        "Type": d[3],
                        "Size": format_size(d[4]),
                        "Notes": d[5],
                        "Uploaded At": d[6]
                    }
                    for d in documents
                ]
            )

        else:

            st.info("No files uploaded to this project yet.")


# =========================================================
# REPORTS
# =========================================================

elif page == "📊 Reports":

    st.title("📊 Reports")

    projects = get_projects()

    col1, col2, col3 = st.columns(3)

    col1.metric("Total Projects", len(projects))

    col2.metric(
        "Active Projects",
        len([p for p in projects if p[3] == "Active"])
    )

    col3.metric("Total Documents", count_documents())


# =========================================================
# AI PAGE
# =========================================================

elif page == "🪄 AI":

    st.title("🪄 AI")

    if is_admin:

        st.write(
            "Analyze project files with AI. "
            "Every result is saved and visible to all users."
        )

    else:

        st.write("AI results saved for each project (read only).")

    projects = get_projects()

    if not projects:

        st.info("No projects yet.")

    else:

        ai_options = {f"{p[0]} — {p[1]}": p[0] for p in projects}

        ai_choice = st.selectbox(
            "Project",
            list(ai_options.keys()),
            key="ai_report_project"
        )

        ai_project = ai_options[ai_choice]

        if is_admin:

            ai_docs = get_documents(ai_project)

            if ai_docs:

                ai_file_options = {
                    f"{d[2]} | {d[6]}": d for d in ai_docs
                }

                ai_file = st.selectbox(
                    "File",
                    list(ai_file_options.keys()),
                    key="ai_report_file"
                )

                ai_prompt = st.text_area(
                    "What should AI do with the file?",
                    value=AI_DEFAULT_INSTRUCTION,
                    height=100,
                    key="ai_report_prompt"
                )

                if st.button("🪄 Run AI", key="ai_report_run"):

                    chosen = ai_file_options[ai_file]

                    with st.spinner("🪄 AI is analyzing..."):

                        file_data = download_bytes(chosen[7])

                        if file_data is None:

                            ai_ok, result = False, "Could not read the file."

                        else:

                            ai_ok, result = analyze_file(
                                chosen[2],
                                file_data,
                                ai_prompt
                            )

                    if ai_ok:

                        save_ai_result(
                            ai_project,
                            chosen[2],
                            ai_prompt,
                            result
                        )

                        st.success("Done. The result is saved below.")

                    else:

                        st.error(result)

            else:

                st.info("No files in this project yet. Upload files first.")

        st.subheader("Previous AI results")

        ai_results = get_ai_results(ai_project)

        if ai_results:

            for i, r in enumerate(ai_results):

                with st.expander(f"{r[0]} — {r[3]}", expanded=(i == 0)):

                    st.caption(f"Request: {r[1]}")

                    st.markdown(r[2])

        else:

            st.info("No AI results for this project yet.")
